# ---------- 1. Outils ----------
import sys, csv
from collections import Counter, defaultdict      # comptages et regroupements
from openpyxl import Workbook                     # création du fichier Excel
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation   # listes déroulantes
from openpyxl.utils import get_column_letter

# ---------- 2. Paramètres (modifiables en ligne de commande) ----------
CSV_ENTREE = sys.argv[1] if len(sys.argv) > 1 else "bulles_lues.csv"
REF_PLAN = sys.argv[2] if len(sys.argv) > 2 else "Q37092-00-IC-DAL-00005"   # numéro du plan lu dans le cartouche
NOM_PDF = sys.argv[3] if len(sys.argv) > 3 else "Plan PID Process 3.pdf"
SORTIE = "liste_IO.xlsx"

# Grille du plan (colonnes 1-12, lignes A-J), mesurée sur l'image à 150 DPI (7021 x 4967 pixels)
X0, LARG_COL = 21, 581.5
Y0, HAUT_LIG = 19, 492.9

VANNES = {"XV", "HV", "AV", "PV", "FV", "LV", "TV"}            # corps de vannes
ACCESSOIRES = {"EV", "ZSH", "ZSL", "ZIT", "PY", "TY", "FY", "LY", "AY"}   # équipements d'une vanne
GRANDEUR = {"L": "niveau", "P": "pression", "T": "température", "F": "débit",
            "A": "analyse", "Z": "position", "D": "densité"}
SIGNAL = {"AI": "4-20 mA", "AO": "4-20 mA", "DI": "TOR (contact)", "DO": "TOR 24 VDC"}   # valeurs par défaut à valider


# ---------- 3. Fonctions utiles ----------
def case_grille(x, y):
    """Convertit une position en pixels en case du plan, par exemple D-10."""
    col = min(12, max(1, int((x - X0) / LARG_COL) + 1))
    lig = "ABCDEFGHIJ"[min(9, max(0, int((y - Y0) / HAUT_LIG)))]
    return f"{lig}-{col}"


def description(f):
    """Description fonctionnelle déduite du code ISA."""
    fixes = {"EV": "Électrovanne", "ZSH": "Fin de course - vanne ouverte", "ZSL": "Fin de course - vanne fermée",
             "ZIT": "Retour de position de vanne", "FSLL": "Contacteur de débit très bas",
             "XV": "Vanne tout-ou-rien", "HV": "Vanne tout-ou-rien", "AV": "Actionneur de vanne",
             "PV": "Vanne de régulation de pression", "FV": "Vanne de régulation de débit",
             "LV": "Vanne de régulation de niveau", "TV": "Vanne de régulation de température"}
    if f in fixes:
        return fixes[f]
    g = GRANDEUR.get(f[0], "?")
    if f.endswith("IT"):
        return f"Transmetteur de {g}"
    if f.endswith("Y"):
        return f"Convertisseur / positionneur ({g})"
    if f.endswith("E"):
        return f"Élément primaire de {g}"
    if f.endswith("I"):
        return f"Indicateur de {g}"
    return f"Instrument ({f})"


def fonction_process(f, io):
    """Valeur de la colonne Process Function : lettre de la grandeur, ou VALVE."""
    if f in ACCESSOIRES or f in VANNES:
        return "VALVE"
    return f[0]


# ---------- 4. Lire le CSV produit par ocr_validation.py ----------
with open(CSV_ENTREE, encoding="utf-8-sig") as fh:
    lignes = list(csv.DictReader(fh, delimiter=";"))

piles = defaultdict(list)
for l in lignes:
    piles[l["pile"]].append(l)
compte_tags = Counter(l["tag"] for l in lignes)      # pour repérer les tags en double

# ---------- 5. Construire les lignes avec niveau de confiance ----------
resultat = []
for numero, p in piles.items():
    fonctions = [l["fonction"] for l in p]
    a_vanne = any(f in VANNES or f in ACCESSOIRES for f in fonctions)
    premieres = {f[0] for f in fonctions}
    coherente = a_vanne or len(premieres) == 1 or len(p) == 1     # PE + LIT dans la même pile = incohérent
    corps = next((f for f in fonctions if f in VANNES), None)
    lettre_boucle = (corps or fonctions[0])[0]                    # H pour HV, L pour LIT...

    for l in p:
        x, y = (int(v) for v in l["pos"].split(","))
        remarques = []
        if l["boucle"] == "???":
            remarques.append("Numéro de boucle illisible")
        if l["statut_fonction"] == "ambigu":
            remarques.append("Fonction ambiguë")
        elif l["statut_fonction"] == "corrigé":
            remarques.append("Fonction corrigée par l'outil")
        if not coherente:
            remarques.append("Fonctions incohérentes dans la pile")
        if compte_tags[l["tag"]] > 1:
            remarques.append("Tag en double")
        if l["confiance"] != "OK":
            remarques.append("Lecture à vérifier")

        if l["boucle"] == "???" or l["statut_fonction"] == "ambigu":
            niveau = "Faible"
        elif remarques:
            niveau = "Moyen"
        else:
            niveau = "Élevé"

        resultat.append(dict(
            tag=l["tag"], f=l["fonction"], io=l["io"], desc=description(l["fonction"]),
            equip=f"{l['zone']}-{corps}-{l['boucle']}" if corps else "",
            loop=f"{l['zone']}-{lettre_boucle}-{l['boucle']}",
            pf=fonction_process(l["fonction"], l["io"]),
            grille=case_grille(x, y), texte=f"{l['zone']} / {l['fonction']} / {l['boucle']}",
            niveau=niveau, remarque="; ".join(remarques)))

avec_io = [r for r in resultat if r["io"]]
sans_io = [r for r in resultat if not r["io"]]

# ---------- 6. Classeur Excel ----------
wb = Workbook()
ws = wb.active
ws.title = "Liste IO"
BLEU = PatternFill("solid", fgColor="1F4E78")
GRIS = PatternFill("solid", fgColor="7F7F7F")
COULEURS = {"Élevé": "C6EFCE", "Moyen": "FFEB9C", "Faible": "FFC7CE"}
GRAS_BLANC = Font(bold=True, color="FFFFFF")

colonnes_cdc = ["I/O Number", "Tag", "Description", "Equipment", "Type", "Signal", "Service",
                "P&ID", "Page", "Datasheet", "Validation", "Commentaire"]
colonnes_plus = ["Confiance", "Position (grille)", "Texte lu", "Loop N°", "Process Function",
                 "Location", "System", "Action signal", "Power supply"]
ws.append(colonnes_cdc + colonnes_plus)
for i, cell in enumerate(ws[1], 1):
    cell.fill = BLEU if i <= len(colonnes_cdc) else GRIS      # bleu = cahier des charges, gris = colonnes supplémentaires
    cell.font = GRAS_BLANC
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

for n, r in enumerate(avec_io, 1):
    ws.append([f"IO-{n:03d}", r["tag"], r["desc"], r["equip"], r["io"], SIGNAL[r["io"]], "",
               REF_PLAN, 1, "", "À valider", r["remarque"],
               r["niveau"], r["grille"], r["texte"], r["loop"], r["pf"],
               "FLD", "DCS", "2 WIRES", "LOOP POWERED (24 VDC)"])
    ws.cell(row=n + 1, column=13).fill = PatternFill("solid", fgColor=COULEURS[r["niveau"]])

# listes déroulantes : Type et Validation
fin = max(len(avec_io) + 1, 2)
dv_type = DataValidation(type="list", formula1='"AI,AO,DI,DO"', allow_blank=True)
dv_val = DataValidation(type="list", formula1='"À valider,Validé,Corrigé,Rejeté"', allow_blank=True)
ws.add_data_validation(dv_type)
ws.add_data_validation(dv_val)
dv_type.add(f"E2:E{fin}")
dv_val.add(f"K2:K{fin}")

ws.freeze_panes = "C2"
ws.auto_filter.ref = ws.dimensions
for i in range(1, ws.max_column + 1):
    ws.column_dimensions[get_column_letter(i)].width = 16
for lettre, largeur in {"B": 20, "C": 34, "D": 20, "L": 44, "O": 22}.items():
    ws.column_dimensions[lettre].width = largeur

# ---------- 7. Feuille des éléments sans signal ----------
ws2 = wb.create_sheet("Sans signal")
ws2.append(["Tag", "Description", "Confiance", "Position (grille)", "Texte lu", "Commentaire"])
for cell in ws2[1]:
    cell.fill, cell.font = BLEU, GRAS_BLANC
for r in sans_io:
    ws2.append([r["tag"], r["desc"], r["niveau"], r["grille"], r["texte"], r["remarque"]])
for i, w in enumerate([20, 34, 12, 16, 22, 44], 1):
    ws2.column_dimensions[get_column_letter(i)].width = w

# ---------- 8. Feuille de résumé ----------
ws3 = wb.create_sheet("Résumé")
ws3.append(["Fichier source", NOM_PDF])
ws3.append(["Référence du plan", REF_PLAN])
ws3.append(["Bulles lues", len(resultat)])
ws3.append(["Lignes d'I/O", len(avec_io)])
ws3.append(["Éléments sans signal", len(sans_io)])
ws3.append([])
ws3.append(["Répartition par type"])
for t in ("AI", "AO", "DI", "DO"):
    ws3.append([t, sum(1 for r in avec_io if r["io"] == t)])
ws3.append([])
ws3.append(["Répartition par confiance"])
for n in ("Élevé", "Moyen", "Faible"):
    ws3.append([n, sum(1 for r in resultat if r["niveau"] == n)])
ws3.append([])
ws3.append(["Note", "Signal, Location, System, Action signal et Power supply sont des valeurs PAR DÉFAUT à valider."])
ws3.append(["Note", "Service, Datasheet et Equipment (hors vannes) demandent les P&ID complets."])
ws3.column_dimensions["A"].width = 28
ws3.column_dimensions["B"].width = 70
for ligne in (7, 12):
    ws3.cell(row=ligne, column=1).font = Font(bold=True)

wb.save(SORTIE)
print(f"{len(avec_io)} lignes d'I/O et {len(sans_io)} éléments sans signal")
print("Confiance :", dict(Counter(r["niveau"] for r in resultat)))
print("Fichier créé :", SORTIE)